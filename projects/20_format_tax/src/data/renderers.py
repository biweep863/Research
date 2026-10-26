"""Render one content tree into every format, and read files back to text.

Text formats (txt, md, tex, html, json) are returned as strings. Binary
formats (docx, pptx, xlsx, pdf) are written to disk and read back the way an
LLM pipeline would see them: either as extracted text (what a loader such as
pypdf or python-docx gives) or as the raw XML inside the zip package.
"""
from __future__ import annotations

import html as _html
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from html.parser import HTMLParser
from pathlib import Path

from .doc_generator import Doc

# formats as the model reads them (Phase R/U); the order is the plotting order
READ_FORMATS = ["txt", "md", "json", "html", "tex", "pdf_text", "docx_text", "pptx_text",
                "xlsx_text", "docx_xml", "pptx_xml", "xlsx_xml"]


def _blocks(doc: Doc):
    """Document order shared by every renderer: (kind, payload, section_idx)."""
    yield "title", doc.title, -1
    yield "byline", f"{doc.author}, {doc.date}", -1
    for i, s in enumerate(doc.sections):
        yield "heading", s.heading, i
        for p in s.paragraphs:
            yield "para", p, i
        if i == 0:
            yield "para", "The yield model used throughout this report is:", i
            yield "equation", doc.equation_text, i
        if i == doc.list_section:
            yield "list", (doc.list_title, doc.list_items), i
        if i == doc.table_section:
            yield "table", (doc.table_caption, doc.table_header, doc.table_rows), i


def _ascii_table(header, rows) -> str:
    w = [max(len(r[j]) for r in [header] + rows) for j in range(len(header))]
    line = lambda r: "  ".join(c.ljust(w[j]) for j, c in enumerate(r)).rstrip()  # noqa: E731
    return "\n".join([line(header), "  ".join("-" * x for x in w)] + [line(r) for r in rows])


def to_txt(doc: Doc) -> str:
    out = []
    for kind, p, _ in _blocks(doc):
        if kind == "title":
            out.append(p.upper())
        elif kind == "heading":
            out.append(f"\n{p}\n")
        elif kind == "list":
            out.append(f"{p[0]}:\n" + "\n".join(f"- {x}" for x in p[1]))
        elif kind == "table":
            out.append(f"{p[0]}:\n" + _ascii_table(p[1], p[2]))
        elif kind == "equation":
            out.append(f"    {p}")
        else:
            out.append(p)
    return "\n".join(out) + "\n"


def to_md(doc: Doc) -> str:
    out = []
    for kind, p, _ in _blocks(doc):
        if kind == "title":
            out.append(f"# {p}\n")
        elif kind == "byline":
            out.append(f"*{p}*\n")
        elif kind == "heading":
            out.append(f"## {p}\n")
        elif kind == "list":
            out.append(f"**{p[0]}**\n\n" + "\n".join(f"- {x}" for x in p[1]) + "\n")
        elif kind == "table":
            cap, h, rows = p
            t = ["| " + " | ".join(h) + " |", "|" + "---|" * len(h)] + ["| " + " | ".join(r) + " |" for r in rows]
            out.append(f"**{cap}**\n\n" + "\n".join(t) + "\n")
        elif kind == "equation":
            out.append(f"$$ {p} $$\n")
        else:
            out.append(p + "\n")
    return "\n".join(out)


def to_json(doc: Doc) -> str:
    d = {"title": doc.title, "author": doc.author, "date": doc.date, "sections": []}
    for kind, p, i in _blocks(doc):
        if kind == "heading":
            d["sections"].append({"heading": p, "content": []})
        elif kind == "para":
            d["sections"][i]["content"].append({"type": "paragraph", "text": p})
        elif kind == "equation":
            d["sections"][i]["content"].append({"type": "equation", "text": p})
        elif kind == "list":
            d["sections"][i]["content"].append({"type": "list", "title": p[0], "items": p[1]})
        elif kind == "table":
            d["sections"][i]["content"].append({"type": "table", "caption": p[0], "header": p[1], "rows": p[2]})
    return json.dumps(d, indent=2, ensure_ascii=False)


def to_html(doc: Doc) -> str:
    e = _html.escape
    out = ["<!DOCTYPE html>", '<html lang="en">', "<head>", '<meta charset="utf-8">',
           f"<title>{e(doc.title)}</title>", "</head>", "<body>"]
    for kind, p, _ in _blocks(doc):
        if kind == "title":
            out.append(f"<h1>{e(p)}</h1>")
        elif kind == "byline":
            out.append(f'<p class="byline"><em>{e(p)}</em></p>')
        elif kind == "heading":
            out.append(f"<h2>{e(p)}</h2>")
        elif kind == "list":
            out.append(f"<p><strong>{e(p[0])}</strong></p>\n<ul>\n" +
                       "\n".join(f"  <li>{e(x)}</li>" for x in p[1]) + "\n</ul>")
        elif kind == "table":
            cap, h, rows = p
            t = [f"<table>\n  <caption>{e(cap)}</caption>", "  <thead>",
                 "    <tr>" + "".join(f"<th>{e(c)}</th>" for c in h) + "</tr>", "  </thead>", "  <tbody>"]
            t += ["    <tr>" + "".join(f"<td>{e(c)}</td>" for c in r) + "</tr>" for r in rows]
            out.append("\n".join(t + ["  </tbody>", "</table>"]))
        elif kind == "equation":
            out.append(f'<p class="equation">{e(p)}</p>')
        else:
            out.append(f"<p>{e(p)}</p>")
    return "\n".join(out + ["</body>", "</html>"]) + "\n"


_TEX_ESC = {"&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_", "{": r"\{", "}": r"\}"}


def _tex(s: str) -> str:
    return "".join(_TEX_ESC.get(c, c) for c in s)


def to_tex(doc: Doc) -> str:
    out = [r"\documentclass[11pt]{article}", r"\usepackage[utf8]{inputenc}", r"\usepackage{amsmath}",
           r"\usepackage{booktabs}", "", rf"\title{{{_tex(doc.title)}}}", rf"\author{{{_tex(doc.author)}}}",
           rf"\date{{{_tex(doc.date)}}}", "", r"\begin{document}", r"\maketitle", ""]
    for kind, p, _ in _blocks(doc):
        if kind in ("title", "byline"):
            continue
        if kind == "heading":
            out.append(rf"\section{{{_tex(p)}}}")
        elif kind == "list":
            out.append(rf"\paragraph{{{_tex(p[0])}}}" + "\n" + r"\begin{itemize}" + "\n" +
                       "\n".join(rf"  \item {_tex(x)}" for x in p[1]) + "\n" + r"\end{itemize}")
        elif kind == "table":
            cap, h, rows = p
            t = [r"\begin{table}[h]", r"\centering", rf"\caption{{{_tex(cap)}}}",
                 r"\begin{tabular}{l" + "r" * (len(h) - 1) + "}", r"\toprule",
                 " & ".join(_tex(c) for c in h) + r" \\", r"\midrule"]
            t += [" & ".join(_tex(c) for c in r) + r" \\" for r in rows]
            out.append("\n".join(t + [r"\bottomrule", r"\end{tabular}", r"\end{table}"]))
        elif kind == "equation":
            out.append(r"\begin{equation}" + "\n  " + doc.equation_latex + "\n" + r"\end{equation}")
        else:
            out.append(_tex(p))
        out.append("")
    return "\n".join(out + [r"\end{document}"]) + "\n"


# ---------------------------------------------------------------- binary writers
def write_docx(doc: Doc, path: Path) -> Path:
    import docx
    d = docx.Document()
    for kind, p, _ in _blocks(doc):
        if kind == "title":
            d.add_heading(p, level=0)
        elif kind == "byline":
            d.add_paragraph(p).runs[0].italic = True
        elif kind == "heading":
            d.add_heading(p, level=1)
        elif kind == "list":
            d.add_paragraph(p[0]).runs[0].bold = True
            for x in p[1]:
                d.add_paragraph(x, style="List Bullet")
        elif kind == "table":
            cap, h, rows = p
            d.add_paragraph(cap, style="Caption")
            t = d.add_table(rows=1 + len(rows), cols=len(h))
            t.style = "Table Grid"
            for i, r in enumerate([h] + rows):
                for j, c in enumerate(r):
                    t.cell(i, j).text = c
        else:
            d.add_paragraph(p)
    d.save(path)
    return path


def write_pptx(doc: Doc, path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = doc.title
    s.placeholders[1].text = f"{doc.author}, {doc.date}"
    for kind, p, _ in _blocks(doc):
        if kind == "heading":
            s = prs.slides.add_slide(prs.slide_layouts[1])
            s.shapes.title.text = p
            body = s.placeholders[1].text_frame
            body.text = ""
            first = True
        elif kind in ("para", "equation"):
            para = body.paragraphs[0] if first else body.add_paragraph()
            para.text, first = p, False
        elif kind == "list":
            ls = prs.slides.add_slide(prs.slide_layouts[1])
            ls.shapes.title.text = p[0]
            tf = ls.placeholders[1].text_frame
            tf.text = p[1][0]
            for x in p[1][1:]:
                tf.add_paragraph().text = x
        elif kind == "table":
            cap, h, rows = p
            ts = prs.slides.add_slide(prs.slide_layouts[5])
            ts.shapes.title.text = cap
            t = ts.shapes.add_table(1 + len(rows), len(h), Inches(0.5), Inches(1.6), Inches(9), Inches(3)).table
            for i, r in enumerate([h] + rows):
                for j, c in enumerate(r):
                    t.cell(i, j).text = c
    prs.save(path)
    return path


def write_xlsx(doc: Doc, path: Path) -> Path:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Content"
    ws.append(["Section", "Type", "Text"])
    ws.append(["", "title", doc.title])
    ws.append(["", "byline", f"{doc.author}, {doc.date}"])
    for kind, p, i in _blocks(doc):
        sec = doc.sections[i].heading if i >= 0 else ""
        if kind in ("para", "equation"):
            ws.append([sec, kind, p])
        elif kind == "list":
            for x in p[1]:
                ws.append([sec, f"list: {p[0]}", x])
        elif kind == "table":
            ws.append([sec, "table", f"{p[0]} (see sheet 'Table')"])
    wt = wb.create_sheet("Table")
    wt.append(doc.table_header)
    for r in doc.table_rows:
        wt.append([r[0]] + [int(c) for c in r[1:]])
    wb.save(path)
    return path


def compile_tex(tex: str, workdir: Path, timeout: float = 60) -> tuple[Path | None, str]:
    """pdflatex in a scratch dir. Returns (pdf path or None, log tail)."""
    workdir.mkdir(parents=True, exist_ok=True)
    (workdir / "doc.tex").write_text(tex, encoding="utf-8")
    try:
        r = subprocess.run(["pdflatex", "-interaction=nonstopmode", "-halt-on-error", "doc.tex"],
                           cwd=workdir, capture_output=True, text=True, timeout=timeout,
                           encoding="utf-8", errors="replace")
    except subprocess.TimeoutExpired:
        return None, "timeout"
    pdf = workdir / "doc.pdf"
    ok = r.returncode == 0 and pdf.exists()
    return (pdf if ok else None), _latex_error(r.stdout)


def _latex_error(log: str) -> str:
    lines = log.splitlines()
    for i, l in enumerate(lines):
        if l.startswith("!"):
            return "\n".join(lines[i:i + 4])
    return "\n".join(lines[-4:])


def html_to_pdf(html: str, path: Path) -> bool:
    from xhtml2pdf import pisa
    with open(path, "wb") as fh:
        res = pisa.CreatePDF(html, dest=fh)
    return not res.err and path.exists() and path.stat().st_size > 0


# ---------------------------------------------------------------- readers
class _Strip(HTMLParser):
    def __init__(self):
        super().__init__()
        self.out, self.skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "head"):
            self.skip += 1
        if tag in ("p", "br", "li", "tr", "h1", "h2", "h3", "caption", "table"):
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "head"):
            self.skip -= 1
        if tag in ("td", "th"):
            self.out.append(" ")

    def handle_data(self, data):
        if not self.skip:
            self.out.append(data)


def html_text(s: str) -> str:
    p = _Strip()
    p.feed(s)
    return re.sub(r"\n\s*\n+", "\n", "".join(p.out)).strip()


def pdf_text(path: Path) -> str:
    from pypdf import PdfReader
    return "\n".join((pg.extract_text() or "") for pg in PdfReader(str(path)).pages)


def docx_text(path: Path) -> str:
    import docx
    from docx.table import Table
    out = []
    for block in docx.Document(str(path)).iter_inner_content():
        if isinstance(block, Table):
            out += [" | ".join(c.text for c in row.cells) for row in block.rows]
        elif block.text.strip():
            out.append(block.text)
    return "\n".join(out)


def pptx_text(path: Path) -> str:
    from pptx import Presentation
    out = []
    for k, slide in enumerate(Presentation(str(path)).slides, 1):
        out.append(f"--- slide {k} ---")
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                out += [" | ".join(c.text for c in row.cells) for row in sh.table.rows]
    return "\n".join(out)


def xlsx_text(path: Path) -> str:
    from openpyxl import load_workbook
    out = []
    for ws in load_workbook(str(path), read_only=True, data_only=True).worksheets:
        out.append(f"# sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            out.append(",".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


def _natkey(n: str):
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", n)]


def zip_xml(path: Path, pattern: str) -> str:
    """Concatenate the XML parts of an OOXML package whose names match `pattern`."""
    with zipfile.ZipFile(path) as z:
        names = sorted((n for n in z.namelist() if re.fullmatch(pattern, n)), key=_natkey)
        return "\n".join(z.read(n).decode("utf-8") for n in names)


XML_PARTS = {"docx": r"word/document\.xml",
             "pptx": r"ppt/slides/slide\d+\.xml",
             "xlsx": r"xl/sharedStrings\.xml|xl/worksheets/sheet\d+\.xml"}

READERS = {"docx": docx_text, "pptx": pptx_text, "xlsx": xlsx_text, "pdf": pdf_text}


def render_all(doc: Doc, workdir: Path) -> dict[str, str]:
    """Every READ_FORMATS string for one doc. Binary files go to `workdir`."""
    workdir.mkdir(parents=True, exist_ok=True)
    out = {"txt": to_txt(doc), "md": to_md(doc), "json": to_json(doc), "html": to_html(doc), "tex": to_tex(doc)}
    for ext, writer in (("docx", write_docx), ("pptx", write_pptx), ("xlsx", write_xlsx)):
        p = writer(doc, workdir / f"{doc.doc_id}.{ext}")
        out[f"{ext}_text"] = READERS[ext](p)
        out[f"{ext}_xml"] = zip_xml(p, XML_PARTS[ext])
    pdf_dir = workdir / f"{doc.doc_id}_tex"
    pdf, log = compile_tex(out["tex"], pdf_dir)
    if pdf is None:
        raise RuntimeError(f"reference LaTeX failed for {doc.doc_id}: {log}")
    shutil.copy(pdf, workdir / f"{doc.doc_id}.pdf")
    out["pdf_text"] = pdf_text(pdf)
    return {k: out[k] for k in READ_FORMATS}


def scratch_dir(prefix: str = "ftax_") -> Path:
    return Path(tempfile.mkdtemp(prefix=prefix))
